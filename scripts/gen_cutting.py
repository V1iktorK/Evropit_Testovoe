# -*- coding: utf-8 -*-
"""
gen_cutting.py — контуры для резки (DXF 1:1) и ведомость заготовок:
 - плоские листовые детали: точный контур (внешний + отверстия) из геометрии;
 - гнутые панели: расчётная заготовка (прямоугольник по площади развёртки);
 - трубы/прутки: распиловка в размер (в ведомости).
"""
import os, sys, math
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import data as D
import gen_drawings as G

import ezdxf
from ezdxf.enums import TextEntityAlignment
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

from OCP.TopExp import TopExp_Explorer
from OCP.TopAbs import TopAbs_FACE, TopAbs_EDGE, TopAbs_WIRE
from OCP.TopoDS import TopoDS
from OCP.BRepAdaptor import BRepAdaptor_Surface, BRepAdaptor_Curve
from OCP.GeomAbs import GeomAbs_Plane, GeomAbs_Line, GeomAbs_Circle, GeomAbs_Ellipse
from OCP.GProp import GProp_GProps
from OCP.BRepGProp import BRepGProp
from OCP.BRepTools import BRepTools_WireExplorer

OUT = os.path.join(D.ROOT, '02_Контуры_резки')
os.makedirs(OUT, exist_ok=True)

def load_shapes():
    st, root = G.load_step()
    G.walk_collect(st, root)
    return G.SHAPES

def face_area(f):
    p = GProp_GProps()
    BRepGProp.SurfaceProperties_s(f, p)
    return p.Mass()

def plane_frame(f):
    ad = BRepAdaptor_Surface(f)
    pl = ad.Plane()
    loc = pl.Location()
    xdir = pl.Position().XDirection()
    ydir = pl.Position().YDirection()
    O = (loc.X(), loc.Y(), loc.Z())
    X = (xdir.X(), xdir.Y(), xdir.Z())
    Y = (ydir.X(), ydir.Y(), ydir.Z())
    return O, X, Y

def to2d(p, frame):
    O, X, Y = frame
    v = (p[0] - O[0], p[1] - O[1], p[2] - O[2])
    return (v[0] * X[0] + v[1] * X[1] + v[2] * X[2],
            v[0] * Y[0] + v[1] * Y[1] + v[2] * Y[2])

def wire_entities(wire, frame):
    """Возвращает список (тип, параметры) 2D-примитивов проволоки."""
    ents = []
    ex = BRepTools_WireExplorer(wire)
    while ex.More():
        e = ex.Current()
        ad = BRepAdaptor_Curve(e)
        t = ad.GetType()
        p0 = ad.Value(ad.FirstParameter())
        p1 = ad.Value(ad.LastParameter())
        a = to2d((p0.X(), p0.Y(), p0.Z()), frame)
        b = to2d((p1.X(), p1.Y(), p1.Z()), frame)
        if t == GeomAbs_Line:
            ents.append(('line', a, b))
        elif t == GeomAbs_Circle:
            c = ad.Circle()
            cc = to2d((c.Location().X(), c.Location().Y(), c.Location().Z()), frame)
            r = c.Radius()
            u1 = math.degrees(ad.FirstParameter())
            u2 = math.degrees(ad.LastParameter())
            if u2 - u1 >= 359.5:
                ents.append(('circle', cc, r))
            else:
                ents.append(('arc', cc, r, u1, u2))
        else:
            # сплайн/эллипс -> полилиния
            from OCP.GCPnts import GCPnts_QuasiUniformDeflection
            disc = GCPnts_QuasiUniformDeflection(ad, 0.02)
            pts = []
            if disc.IsDone():
                for i in range(1, disc.NbPoints() + 1):
                    p = disc.Value(i)
                    pts.append(to2d((p.X(), p.Y(), p.Z()), frame))
            if len(pts) >= 2:
                ents.append(('poly', pts))
        ex.Next()
    return ents

def extract_contour(shape):
    """Внешний и внутренние контуры плоской детали (по самой большой грани)."""
    best, best_area = None, 0.0
    ex = TopExp_Explorer(shape, TopAbs_FACE)
    while ex.More():
        f = TopoDS.Face_s(ex.Current())
        ad = BRepAdaptor_Surface(f)
        if ad.GetType() == GeomAbs_Plane:
            a = face_area(f)
            if a > best_area:
                best, best_area = f, a
        ex.Next()
    if best is None:
        return None
    frame = plane_frame(best)
    wires = []
    wex = TopExp_Explorer(best, TopAbs_WIRE)
    while wex.More():
        w = TopoDS.Wire_s(wex.Current())
        wires.append(w)
        wex.Next()
    if not wires:
        return None
    # внешний контур — проволока с наибольшим габаритом
    def wire_size(w):
        ents = wire_entities(w, frame)
        xs, ys = [], []
        for e in ents:
            if e[0] == 'line':
                xs += [e[1][0], e[2][0]]; ys += [e[1][1], e[2][1]]
            elif e[0] == 'circle':
                xs += [e[1][0] - e[2], e[1][0] + e[2]]; ys += [e[1][1] - e[2], e[1][1] + e[2]]
            elif e[0] == 'arc':
                xs += [e[1][0] - e[2], e[1][0] + e[2]]; ys += [e[1][1] - e[2], e[1][1] + e[2]]
            elif e[0] == 'poly':
                xs += [p[0] for p in e[1]]; ys += [p[1] for p in e[1]]
        if not xs:
            return 0, ents
        return (max(xs) - min(xs)) * (max(ys) - min(ys)), ents
    sized = sorted(((wire_size(w), w) for w in wires), key=lambda t: -t[0][0])
    outer = sized[0][0][1]
    inner = [t[0][1] for t in sized[1:]]
    return outer, inner

def save_cut_dxf(path, name, outer, inner, note_lines, qty, material):
    doc = ezdxf.new('R2018', setup=True)
    doc.header['$INSUNITS'] = 4
    doc.styles.add('GOST', font='Arial')
    msp = doc.modelspace()
    doc.layers.add('CUT', color=1)
    doc.layers.add('CUT_INNER', color=3)
    doc.layers.add('MARK', color=5)

    def put(ents, layer):
        for e in ents:
            if e[0] == 'line':
                msp.add_line(e[1], e[2], dxfattribs={'layer': layer})
            elif e[0] == 'circle':
                msp.add_circle(e[1], e[2], dxfattribs={'layer': layer})
            elif e[0] == 'arc':
                msp.add_arc(e[1], e[2], e[3], e[4], dxfattribs={'layer': layer})
            elif e[0] == 'poly':
                msp.add_lwpolyline(e[1], dxfattribs={'layer': layer})
    put(outer, 'CUT')
    for w in inner:
        put(w, 'CUT_INNER')
    # маркировка: позиционное обозначение + количество
    xs, ys = [], []
    for e in outer:
        if e[0] == 'line':
            xs += [e[1][0], e[2][0]]; ys += [e[1][1], e[2][1]]
        elif e[0] in ('circle', 'arc'):
            xs += [e[1][0] - e[2], e[1][0] + e[2]]; ys += [e[1][1] - e[2], e[1][1] + e[2]]
    if xs:
        x0 = min(xs)
        y0 = min(ys) - 12
        msp.add_text(f"{name}  x{qty}  ({material})",
                     dxfattribs={'style': 'GOST', 'height': 6, 'layer': 'MARK'}
                     ).set_placement((x0, y0), align=TextEntityAlignment.LEFT)
    doc.saveas(path)

# ------------------------------------------------------------------ main
def main():
    shapes = load_shapes()
    rows = []
    for name in sorted(D.TOTALS):
        if not D.is_custom(name) or D.is_purchased(name) or name in D.TREE:
            continue
        kind = D.part_kind(name)
        g = D.geom_of(name)
        mass = D.part_mass_kg(name)
        qty = D.TOTALS[name]
        mat = D.part_material_ru(name).split(' (')[0]
        dims = sorted([g['dx'], g['dy'], g['dz']])
        th, mid, big = dims
        if kind == 'plate':
            try:
                cont = extract_contour(shapes.get(name) or G.SHAPES.get(name))
            except Exception:
                cont = None
            if cont:
                outer, inner = cont
                fn = f"{name}__раскрой.dxf"
                save_cut_dxf(os.path.join(OUT, fn), name, outer, inner, [], qty, mat)
                rows.append([name, D.part_name(name), mat, 'Лист / плоская деталь',
                             f"{fmt(big)}×{fmt(mid)}, s={fmt(th)}", qty, round(mass * qty, 2),
                             f"Резка лазерная/плазменная по DXF ({fn})"])
            else:
                rows.append([name, D.part_name(name), mat, 'Лист / плоская деталь',
                             f"{fmt(big)}×{fmt(mid)}, s={fmt(th)}", qty, round(mass * qty, 2),
                             'Контур не извлечён — уточнить'])
        elif kind == 'bent':
            t_est = D.sheet_thickness_est(name)
            A_flat = g['volume'] / max(t_est, 0.1)     # мм2 срединной поверхности
            L_b = big * 1.04                            # припуск на гибы/полки
            W_b = A_flat / L_b
            fn = f"{name}__заготовка.dxf"
            outer = [('line', (0, 0), (L_b, 0)), ('line', (L_b, 0), (L_b, W_b)),
                     ('line', (L_b, W_b), (0, W_b)), ('line', (0, W_b), (0, 0))]
            save_cut_dxf(os.path.join(OUT, fn), name, outer, [], [], qty, mat)
            rows.append([name, D.part_name(name), mat, 'Лист гнутый (заготовка)',
                         f"{fmt(L_b, 0)}×{fmt(W_b, 0)}, s={fmt(t_est)}", qty,
                         round(mass * qty, 2),
                         f"Раскрой по DXF ({fn}); гибка по чертежу, размеры заготовки уточнить по КД гибки"])
        elif kind in ('tube', 'profile'):
            t_est = D.sheet_thickness_est(name)
            rows.append([name, D.part_name(name), mat, 'Труба/профиль',
                         f"{fmt(mid)}×{fmt(th)}, L={fmt(big)}", qty, round(mass * qty, 2),
                         'Распиловка в размер на ленточнопильном станке'])
        elif kind == 'shaft':
            rows.append([name, D.part_name(name), mat, 'Пруток',
                         f"Ø{fmt(mid)}, L={fmt(big)}", qty, round(mass * qty, 2),
                         'Распиловка в размер; дальнейшая обработка по чертежу'])
        else:
            rows.append([name, D.part_name(name), mat, 'Заготовка (уточнить)',
                         f"{fmt(big)}×{fmt(mid)}×{fmt(th)}", qty, round(mass * qty, 2),
                         'Заготовка по технологу (сварная/литая), см. чертёж'])

    # ведомость
    wb = Workbook()
    ws = wb.active
    ws.title = 'Ведомость заготовок'
    hdr = ['№', 'Обозначение', 'Наименование', 'Материал (РФ)', 'Вид заготовки',
           'Размер заготовки, мм', 'Кол.', 'Масса, кг', 'Операция / технология']
    widths = [5, 24, 38, 30, 24, 30, 6, 12, 60]
    thin = Side(style='thin')
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    for i, (h, w) in enumerate(zip(hdr, widths), 1):
        c = ws.cell(1, i, h)
        c.font = Font(bold=True, size=10)
        c.fill = PatternFill('solid', fgColor='D9E1F2')
        c.border = bd
        c.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    for r, row in enumerate(rows, 2):
        for i, v in enumerate([r - 1] + row, 1):
            c = ws.cell(r, i, v)
            c.font = Font(size=10)
            c.border = bd
            c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.cell(len(rows) + 2, 8, round(sum(x[6] for x in rows), 1)).font = Font(bold=True, size=10)
    ws.cell(len(rows) + 2, 3, 'ИТОГО масса заготовок, кг').font = Font(bold=True, size=10)
    wb.save(os.path.join(OUT, 'Ведомость_заготовок.xlsx'))
    print('rows:', len(rows))
    print('saved', OUT)

def fmt(v, nd=1):
    s = f"{v:.{nd}f}".rstrip('0').rstrip('.')
    return s.replace('.', ',')

if __name__ == '__main__':
    main()
