# -*- coding: utf-8 -*-
"""
gen_bom.py — спецификация (ГОСТ 2.106) + сводный BOM с российскими аналогами.
Выход: 03_BOM/BOM_PSH24-000.00.xlsx, BOM_PSH24-000.00.csv
"""
import os, sys, csv
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import data as D
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

OUT = os.path.normpath(os.path.join(D.ROOT, '03_BOM'))
os.makedirs(OUT, exist_ok=True)

thin = Side(style='thin')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR_FILL = PatternFill('solid', fgColor='D9E1F2')
GRP_FILL = PatternFill('solid', fgColor='F2F2F2')
BOLD = Font(bold=True, size=10)
NORM = Font(size=10)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)

ROOT = 'PSH24-000.00'

def kind_ru(kind):
    return {'plate': 'листовая', 'bent': 'листовая гнутая', 'tube': 'профильная',
            'shaft': 'стержневая', 'profile': 'профильная', 'other': 'механическая',
            'purchased': 'покупное'}.get(kind, '')

def spec_rows(node, depth=0, out=None, path=()):
    """Рекурсивная спецификация: строки (depth, name, qty)."""
    if out is None:
        out = []
    kids = D.TREE.get(node, {})
    for ch, q in sorted(kids.items()):
        out.append((depth, ch, q))
        if ch in D.TREE and ch not in path:
            spec_rows(ch, depth + 1, out, path + (node,))
    return out

# ----------------------------------------------------------------- Excel
wb = Workbook()

# ===== Лист 1: спецификация (иерархия) =====
ws = wb.active
ws.title = 'Спецификация'
headers = ['Поз.', 'Наименование / Обозначение', 'Тип', 'Кол.',
           'Материал (РФ)', 'Масса ед., кг', 'Масса общ., кг',
           'Российский аналог / ГОСТ', 'Примечание']
widths = [6, 42, 16, 6, 40, 12, 12, 46, 30]
for i, (h, w) in enumerate(zip(headers, widths), 1):
    c = ws.cell(1, i, h)
    c.font = BOLD
    c.fill = HDR_FILL
    c.border = BORDER
    c.alignment = CENTER
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = 'A2'

def row_fill(ws, r, vals, fill=None, bold=False):
    for i, v in enumerate(vals, 1):
        c = ws.cell(r, i, v)
        c.font = BOLD if bold else NORM
        c.border = BORDER
        c.alignment = WRAP if i in (2, 5, 8, 9) else CENTER
        if fill:
            c.fill = fill

r = 2
row_fill(ws, r, ['', ROOT + ' — Бункер (изделие), сборочная единица верхнего уровня',
                 'сборка', 1, '', '', round(D.part_mass_kg(ROOT), 1) if ROOT in D.GEOM else '',
                 '', ''], GRP_FILL, True)
r += 1
positions = {}
for depth, name, qty in spec_rows(ROOT):
    positions[name] = positions.get(name, 0) + 1
    is_asm = name in D.TREE
    kind = 'сборка' if is_asm else D.part_kind(name)
    mass = D.part_mass_kg(name)
    row_fill(ws, r, [
        f"{depth}",                       # уровень вложенности
        ('    ' * depth) + f"{name} — {D.part_name(name)}",
        kind, qty,
        D.part_material_ru(name),
        round(mass, 2), round(mass * qty, 2),
        D.analog(name) if D.is_purchased(name) else 'Изготовить по КД PSH24',
        'входит в сборку' if is_asm else ('покупное изделие' if D.is_purchased(name) else kind_ru(D.part_kind(name))),
    ], GRP_FILL if is_asm else None)
    r += 1

# ===== Лист 2: сводный плоский BOM =====
ws2 = wb.create_sheet('BOM (сводный)')
headers2 = ['№', 'Обозначение', 'Наименование', 'Кол. (всего)',
            'Материал (EN)', 'Материал (РФ / ГОСТ)', 'Масса ед., кг',
            'Масса общ., кг', 'Российский аналог / ГОСТ', 'Тип', 'Чертёж']
for i, (h, w) in enumerate(zip(headers2, [5, 26, 40, 10, 14, 40, 11, 12, 48, 16, 12]), 1):
    c = ws2.cell(1, i, h)
    c.font = BOLD; c.fill = HDR_FILL; c.border = BORDER; c.alignment = CENTER
    ws2.column_dimensions[get_column_letter(i)].width = w
ws2.freeze_panes = 'A2'

r = 2
flat = sorted(D.TOTALS.items(), key=lambda kv: (not D.is_custom(kv[0]), kv[0]))
for i, (name, qty) in enumerate(flat, 1):
    mass = D.part_mass_kg(name)
    drawing = 'да' if D.is_custom(name) else '—'
    row_fill(ws2, r, [i, name, D.part_name(name), qty,
                      D.MAT_EN.get(name, ''), D.part_material_ru(name),
                      round(mass, 3), round(mass * qty, 2),
                      D.analog(name), D.part_kind(name), drawing])
    r += 1
# итог
row_fill(ws2, r, ['', 'ИТОГО масса изделия, кг:', '', '', '', '', '',
                  round(sum(D.part_mass_kg(n) * q for n, q in D.TOTALS.items()), 1),
                  '', '', ''], HDR_FILL, True)

# ===== Лист 3: материалы =====
ws3 = wb.create_sheet('Материалы (свод)')
for i, (h, w) in enumerate(zip(['Материал (EN)', 'Материал (РФ / ГОСТ)', 'Масса, кг', 'Позиций'],
                               [18, 52, 14, 10]), 1):
    c = ws3.cell(1, i, h); c.font = BOLD; c.fill = HDR_FILL; c.border = BORDER; c.alignment = CENTER
    ws3.column_dimensions[get_column_letter(i)].width = w
agg = {}
for name, qty in D.TOTALS.items():
    m = D.MAT_EN.get(name)
    mass = D.part_mass_kg(name) * qty
    a = agg.setdefault(m, [0.0, 0])
    a[0] += mass
    a[1] += 1
r = 2
for m, (mass, cnt) in sorted(agg.items(), key=lambda kv: -kv[1][0]):
    row_fill(ws3, r, [m, D.MAT_RU.get(m, ''), round(mass, 1), cnt])
    r += 1

# ===== Лист 4: аналоги =====
ws4 = wb.create_sheet('Аналоги ГОСТ')
for i, (h, w) in enumerate(zip(['Обозначение (модель)', 'Наименование', 'Российский аналог / ГОСТ'],
                               [34, 44, 70]), 1):
    c = ws4.cell(1, i, h); c.font = BOLD; c.fill = HDR_FILL; c.border = BORDER; c.alignment = CENTER
    ws4.column_dimensions[get_column_letter(i)].width = w
r = 2
for name in sorted(D.STD):
    if name in D.TOTALS:
        row_fill(ws4, r, [name, D.STD[name][0], D.STD[name][1]])
        r += 1

xlsx_path = os.path.join(OUT, 'BOM_PSH24-000.00.xlsx')
wb.save(xlsx_path)
print('saved', xlsx_path)

# ----------------------------------------------------------------- CSV
csv_path = os.path.join(OUT, 'BOM_PSH24-000.00.csv')
with open(csv_path, 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.writer(f, delimiter=';')
    w.writerow(['Обозначение', 'Наименование', 'Кол', 'Материал EN', 'Материал РФ',
                'Масса ед кг', 'Масса общ кг', 'Аналог ГОСТ', 'Тип'])
    for name, qty in flat:
        mass = D.part_mass_kg(name)
        w.writerow([name, D.part_name(name), qty, D.MAT_EN.get(name, ''),
                    D.part_material_ru(name), f"{mass:.3f}", f"{mass*qty:.2f}",
                    D.analog(name), D.part_kind(name)])
print('saved', csv_path)
